"""Generate defense_presentation.pptx + OCR QA on slide previews."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DOCS = Path(__file__).resolve().parent
NUMBERS = DOCS / "defense_verified_numbers.json"
OUT = DOCS / "defense_presentation.pptx"
OUT_FALLBACK = DOCS / "defense_presentation_styled.pptx"
ASSETS = DOCS / "defense_assets"
PREVIEW = ASSETS / "preview"

SLIDE_W = Inches(10)
SLIDE_H = Inches(7.5)
ML = 0.5
CONTENT_W = 9.0
FOOTER_Y = 7.05

C_TEXT = RGBColor(30, 30, 30)
C_MUTED = RGBColor(90, 90, 90)
C_LINE = RGBColor(180, 180, 180)
C_HDR = RGBColor(245, 245, 245)
C_ACCENT = RGBColor(0, 70, 130)


def load_numbers() -> dict:
    with NUMBERS.open(encoding="utf-8") as f:
        return json.load(f)


class Deck:
    def __init__(self, prs: Presentation) -> None:
        self.prs = prs
        self.n = 0

    def slide(self):
        self.n += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        return s

    def title_line(self, slide, text: str, *, y: float = 0.35) -> float:
        tb = slide.shapes.add_textbox(Inches(ML), Inches(y), Inches(CONTENT_W), Inches(0.55))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = C_TEXT
        p.font.name = "Calibri"
        line = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(ML), Inches(y + 0.55), Inches(CONTENT_W), Inches(0.012)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = C_LINE
        line.line.fill.background()
        return y + 0.68

    def footer(self, slide) -> None:
        tb = slide.shapes.add_textbox(Inches(ML), Inches(FOOTER_Y), Inches(7), Inches(0.3))
        tb.text_frame.paragraphs[0].text = "RAG · годовой проект ВШЭ 2025–2026 · Серов А.А."
        tb.text_frame.paragraphs[0].font.size = Pt(9)
        tb.text_frame.paragraphs[0].font.color.rgb = C_MUTED
        num = slide.shapes.add_textbox(Inches(9.2), Inches(FOOTER_Y), Inches(0.5), Inches(0.3))
        num.text_frame.paragraphs[0].text = str(self.n)
        num.text_frame.paragraphs[0].font.size = Pt(9)
        num.text_frame.paragraphs[0].font.color.rgb = C_MUTED
        num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    def bullets(
        self,
        slide,
        items: list[str],
        *,
        top: float,
        left: float = ML,
        width: float = CONTENT_W,
        size: int = 14,
        bold_first: bool = False,
    ) -> None:
        tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(FOOTER_Y - top - 0.12))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(size)
            p.font.color.rgb = C_TEXT
            p.font.name = "Calibri"
            p.font.bold = bold_first and i == 0
            p.space_after = Pt(7)

    def note(self, slide, text: str, *, top: float, left: float = ML, width: float = CONTENT_W) -> None:
        tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(10)
        p.font.color.rgb = C_MUTED
        p.font.italic = True

    def table(
        self,
        slide,
        headers: list[str],
        rows: list[list[str]],
        *,
        top: float,
        left: float = ML,
        width: float = CONTENT_W,
        col_widths: list[float],
        font_size: int = 10,
    ) -> None:
        rh = 0.34
        shape = slide.shapes.add_table(
            len(rows) + 1, len(headers), Inches(left), Inches(top), Inches(width), Inches(rh * (len(rows) + 1))
        )
        tbl = shape.table
        for j, w in enumerate(col_widths):
            tbl.columns[j].width = Inches(w)
        for j, h in enumerate(headers):
            c = tbl.cell(0, j)
            c.text = h
            c.fill.solid()
            c.fill.fore_color.rgb = C_HDR
            for p in c.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(font_size)
                p.font.color.rgb = C_TEXT
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                c = tbl.cell(i, j)
                c.text = val
                for p in c.text_frame.paragraphs:
                    p.font.size = Pt(font_size)
                    p.font.color.rgb = C_TEXT
                c.vertical_anchor = MSO_ANCHOR.MIDDLE

    def picture_fit(
        self,
        slide,
        filename: str,
        *,
        top: float,
        left: float = ML,
        max_w: float | None = None,
        max_h: float | None = None,
    ) -> float:
        path = ASSETS / filename
        if not path.is_file():
            tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(4), Inches(0.35))
            tb.text_frame.paragraphs[0].text = f"[нет: {filename}]"
            return top + 0.35

        mw = max_w if max_w is not None else CONTENT_W
        mh = max_h if max_h is not None else FOOTER_Y - top - 0.15

        with Image.open(path) as im:
            px_w, px_h = im.size
        aspect = px_w / px_h
        w_by_width = mw
        h_by_width = w_by_width / aspect
        if h_by_width <= mh:
            w, h = w_by_width, h_by_width
        else:
            h = mh
            w = h * aspect

        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(w))
        return top + h


def build() -> Path:
    n = load_numbers()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    d = Deck(prs)

    chunks = f"{n['corpus']['chunks']:,}".replace(",", " ")
    docs = n["eda"]["documents"]
    gold = n["corpus"]["gold_questions"]
    max_r10 = n["full_corpus_substring_recall_at_10"]["max_pct"]
    prod_r10 = n["production_recommendation"]["retrieval_r_at_10_pct"]
    prod_emb = n["production_recommendation"]["embedder"]
    prod_rr = n["production_recommendation"]["reranker"]
    ans_sim = n["answer_eval"]["best_avg_ref_sim"]
    ans_cfg = n["answer_eval"]["best_config"]
    squad = n["validity"]["squad_sanity_r_at_10_pct"]
    neg_h = n["validity"]["neg_control_honest_pct"]
    neg_s = n["validity"]["neg_control_shuffled_pct"]
    class_a = n["validity"]["metrics_ladder_class_a_pct_approx"]
    fetch_eval = n["full_corpus_substring_recall_at_10"]["fetch_k_eval"]
    fetch_prod = n["service_code_defaults"]["retriever_fetch_k"]
    top_k_prod = n["service_code_defaults"]["retriever_k"]

    # ── 1 Титул ──────────────────────────────────────────────────────────
    s = d.slide()
    tb = s.shapes.add_textbox(Inches(ML), Inches(1.45), Inches(CONTENT_W), Inches(1.6))
    tb.text_frame.paragraphs[0].text = "Применение LLM для извлечения\nинформации из текстов"
    tb.text_frame.paragraphs[0].font.size = Pt(30)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = C_TEXT
    sub = s.shapes.add_textbox(Inches(ML), Inches(3.2), Inches(CONTENT_W), Inches(2.8))
    sf = sub.text_frame
    sf.word_wrap = True
    for i, line in enumerate(
        [
            "RAG-сервис + воспроизводимый eval на full corpus (61k чанков)",
            "Серов Арсений Алексеевич · куратор: Малюшитский Кирилл",
            "GitHub: Haradrick228/LLM_EXTRACT_INFO_FROM_TEXT_52 · 2025–2026",
        ]
    ):
        sp = sf.paragraphs[0] if i == 0 else sf.add_paragraph()
        sp.text = line
        sp.font.size = Pt(14 if i else 15)
        sp.font.color.rgb = C_MUTED if i else C_ACCENT
        sp.space_after = Pt(6)
    d.footer(s)

    # ── 2 Задача (без цифр корпуса — они на EDA) ─────────────────────────
    s = d.slide()
    y = d.title_line(s, "Задача и границы проекта")
    d.bullets(
        s,
        [
            "Q&A по образовательным материалам: пользователь задаёт вопрос — система ищет чанки и генерирует ответ",
            "Deliverables: Telegram-бот · REST POST /chat :8001 · ML-service :8000 · HTML-отчёты eval",
            "В scope: dense RAG, cross-encoder rerank, frozen eval protocol, observability (traces, Postgres log)",
            "Вне scope (отложено): NER/IE сущностей, GraphRAG (stub), F2 hybrid как prod-конфиг",
            "Критерий успеха: работающий сервис + доказуемая методология (6 чекпоинтов CP2–CP7)",
        ],
        top=y + 0.12,
        size=14,
    )
    d.footer(s)

    # ── 3 EDA: интерпретация (цифры — на скринах, не в буллетах) ─────────
    s = d.slide()
    y = d.title_line(s, "Данные и подготовка корпуса (CP2)")
    d.bullets(
        s,
        [
            "Pipeline: загрузка → очистка → SHA-256 → TF-IDF → chunking → индексация в Chroma",
            "RU-dominant корпус → multilingual embedders (mpnet, BGE-M3, FRIDA, E5, Jina, USER-bge-m3)",
            "Phase B: chunk 650/250 — лучший компромисс; re-chunking на full corpus ухудшал recall",
            "Артефакты CP2: eda_dashboard.html, embedding analysis (оценка куратора 0.5)",
        ],
        top=y + 0.08,
        size=13,
    )
    img_top = y + 1.55
    half_w = (CONTENT_W - 0.2) / 2
    d.picture_fit(s, "eda_key_figures.png", top=img_top, left=ML, max_w=half_w, max_h=FOOTER_Y - img_top - 0.25)
    d.picture_fit(s, "eda_quality.png", top=img_top, left=ML + half_w + 0.2, max_w=half_w, max_h=FOOTER_Y - img_top - 0.25)
    d.footer(s)

    # ── 4 Архитектура: prod vs eval (таблица, без повтора методологии) ──
    s = d.slide()
    y = d.title_line(s, "Архитектура: prod vs eval")
    d.table(
        s,
        ["Слой", "Production", "Eval / research"],
        [
            ["Клиент", "Telegram-бот, REST :8001", "replay_experiment, run_profile_metrics"],
            ["Retrieval", f"Chroma USER-bge-m3, fetch_k={fetch_prod}, top_k={top_k_prod}", f"7 коллекций Chroma, fetch_k={fetch_eval}"],
            ["Rerank", "BGE reranker (RAG_PROFILE)", "MiniLM / BGE / Jina — матрица 7×3"],
            ["LLM", "DeepSeek / Ollama", "DeepSeek + answer-eval (avg_ref_sim)"],
            ["Конфиг", "RAG_PROFILE → rag_profiles.yaml", "final_research_protocol.yaml (frozen)"],
            ["Логи", "Postgres request_log", "trace.jsonl, HTML reports (CP7)"],
        ],
        top=y + 0.12,
        col_widths=[1.4, 3.8, 3.8],
        font_size=10,
    )
    d.note(
        s,
        "Research-модули (не prod): BM25, Hybrid RRF, HyDE, BGE-M3 multi-vector — core/rag/",
        top=y + 2.55,
    )
    d.footer(s)

    # ── 5 Методология (КАК мерим, без итоговых 6.4%) ─────────────────────
    s = d.slide()
    y = d.title_line(s, "Методология оценки")
    d.bullets(
        s,
        [
            f"Gold: {gold} вопросов (eval_set_gold_v2), поля references + answer_ref; SHA256 зафиксирован в protocol",
            "Primary gate — substring::recall@10: эталонная фраза ∈ текст чанка в top-k (жёсткий критерий для отчёта)",
            "Диагностика (другая шкала!): token_f1, semantic hit (cosine ≥ 0.62) — не смешивать с substring %",
            "Answer layer (отдельный прогон): avg_ref_sim = cosine(embed(ответ LLM), embed(answer_ref))",
            f"Metrics ladder: ~{class_a}% class A — эталон substring отсутствует в {chunks} чанках → потолок recall",
            f"Validity: SQuAD sanity {squad}% · neg-control {neg_h}% → {neg_s}% (shuffle) · 2 121 retrieval-оценок",
        ],
        top=y + 0.1,
        size=13,
    )
    d.footer(s)

    # ── 6 Результаты: сводная таблица (без скрина — интерпретация) ─────
    s = d.slide()
    y = d.title_line(s, "Результаты: что выбрали и почему")
    d.table(
        s,
        ["Конфигурация", "R@10", "avg_ref_sim", "Роль"],
        [
            [f"{prod_emb} + {prod_rr}", f"{prod_r10}%", "0.642", "Prod (скорость индекса + rerank)"],
            ["full_frida + rerank_bge_v2_m3", f"{max_r10}%", "0.650", "Max substring retrieval"],
            [ans_cfg.replace(" + ", " + "), f"{n['answer_eval']['retrieval_r_at_10_pct']}%", f"{ans_sim:.3f}", "Max качество ответа LLM"],
            ["full_mpnet + MiniLM", "0.5–2.0%", "—", "Baseline Phase B (старт ~2%)"],
        ],
        top=y + 0.1,
        col_widths=[3.6, 0.9, 1.1, 3.4],
        font_size=10,
    )
    d.bullets(
        s,
        [
            f"Матрица 7×3 на full corpus: max substring R@10 = {max_r10}% — упирается в class A (~{class_a}%), не в баг pipeline",
            f"Semantic recall@10 до {n['validity']['semantic_r_at_10_max_pct']}% — retrieval «по смыслу» существенно выше substring",
            "⚠ Не путать: dashboard 64% и fixture rerank 95.5% — eval на ~101 chunk, не на 61k (исправлено 31.05.2026)",
            f"Канонический run: {n['full_corpus_substring_recall_at_10']['canonical_run']}",
        ],
        top=y + 1.85,
        size=12,
    )
    d.footer(s)

    # ── 7 Матрица: только скрин (детали — на слайде выше) ───────────────
    s = d.slide()
    y = d.title_line(s, "Матрица 7 embedder × 3 reranker — substring Recall@10")
    d.note(
        s,
        "Полная таблица · fetch_k=40 · 61 543 чанка · отчёт: final_canonical_experiment_2026_05.html",
        top=y,
    )
    d.picture_fit(s, "canonical_results.png", top=y + 0.28, max_h=FOOTER_Y - y - 0.35)
    d.footer(s)

    # ── 8 Чекпоинты: уникальный результат каждого CP ─────────────────────
    s = d.slide()
    y = d.title_line(s, "Чекпоинты: что сделано на каждом этапе")
    d.table(
        s,
        ["CP", "Этап", "Ключевой результат", "Артефакт"],
        [
            ["CP2", "EDA", f"{docs} док., ru {n['eda']['dominant_language_ru_pct']}%, дедуп {n['eda']['duplicate_rate_pct']}%", "eda_dashboard.html"],
            ["CP3", "Baseline RAG", "Gold 101 Q, mpnet+MiniLM ~2% R@10 на full corpus", "eval_set_gold_v2.jsonl"],
            ["CP5", "Retrieve Phase B", "6 embedders на fixture → 7 на full corpus, chunk 650/250", "phase_b collections"],
            ["CP4", "Сервис", "TG + REST + ML-service, пром. защита «все 9»", "bots/telegram, bots/rest_api"],
            ["CP6", "DL / eval", f"Матрица 7×3, max {max_r10}% R@10, frozen protocol", "FINAL_CANONICAL_REPORT"],
            ["CP7", "Observability", "Traces, error taxonomy, neg-control, metrics ladder", "rag_observability report"],
        ],
        top=y + 0.08,
        col_widths=[0.55, 1.3, 4.0, 3.15],
        font_size=9,
    )
    d.footer(s)

    # ── 9 Демо: что показать на защите ──────────────────────────────────
    s = d.slide()
    y = d.title_line(s, "Демонстрация на защите")
    d.bullets(
        s,
        [
            "Telegram: профильный вопрос по материалам курса → ответ с опорой на найденные чанки",
            "Off-topic / нет в базе → «Недостаточно данных в базе знаний» (не галлюцинация)",
            f"RAG_PROFILE={prod_emb} + {prod_rr} — переключение профиля без правки кода",
            "REST: POST /chat :8001 (JSON) · ML-service: POST /forward :8000",
            "CP7: показать trace.jsonl / HTML observability — где retrieval промахнулся (class A vs D)",
        ],
        top=y + 0.12,
        size=14,
    )
    d.note(
        s,
        "Feedback/NPS, загрузка документов в библиотеку — доп. функции бота",
        top=6.35,
    )
    d.footer(s)

    # ── 10 Выводы: инсайты, не повтор цифр ──────────────────────────────
    s = d.slide()
    y = d.title_line(s, "Выводы для комиссии")
    d.bullets(
        s,
        [
            "Embedder и reranker — главные рычаги: mpnet→FRIDA/USER-bge-m3 и MiniLM→BGE дают кратный прирост",
            "Chunking 650/250 достаточен; дальнейший re-chunking не улучшал substring recall",
            f"Низкий substring recall ({max_r10}%) объясняется разметкой gold (~{class_a}% class A), а не поломкой пайплайна",
            f"LLM-компонент: avg_ref_sim {ans_sim} при imperfect retrieval — ответ может быть полезен без exact chunk hit",
            "Воспроизводимость: frozen protocol, replay_experiment, единый canonical report v2 (2026-06-02)",
        ],
        top=y + 0.1,
        size=14,
    )
    thanks = s.shapes.add_textbox(Inches(ML), Inches(6.5), Inches(CONTENT_W), Inches(0.4))
    thanks.text_frame.paragraphs[0].text = "Спасибо за внимание · готов показать бота и отчёт"
    thanks.text_frame.paragraphs[0].font.size = Pt(13)
    thanks.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    d.footer(s)

    try:
        prs.save(OUT)
        return OUT
    except PermissionError:
        prs.save(OUT_FALLBACK)
        return OUT_FALLBACK


def export_slide_previews(pptx_path: Path) -> list[Path]:
    PREVIEW.mkdir(parents=True, exist_ok=True)
    for old in PREVIEW.glob("slide_*.png"):
        old.unlink()
    import win32com.client  # type: ignore

    app = win32com.client.Dispatch("PowerPoint.Application")
    app.Visible = 1
    pres = app.Presentations.Open(str(pptx_path.resolve()), WithWindow=False)
    out: list[Path] = []
    try:
        for i in range(1, pres.Slides.Count + 1):
            p = PREVIEW / f"slide_{i:02d}.png"
            pres.Slides(i).Export(str(p.resolve()), "PNG", 1920, 1080)
            out.append(p)
    finally:
        pres.Close()
        app.Quit()
    return out


def ocr_qa_previews(preview_dir: Path) -> list[dict]:
    """OCR each slide preview; check readability and expected keywords."""
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

    expectations: dict[int, list[str]] = {
        1: ["LLM", "RAG", "GitHub", "corpus"],
        2: ["Deliverables", "Telegram", "GraphRAG", "CP"],
        3: ["Pipeline", "Chroma", "650/250", "eda_dashboard"],
        4: ["Production", "Eval", "fetch_k", "Chroma"],
        5: ["Gold", "substring", "SQuAD", "class A"],
        6: ["R@10", "Prod", "frida", "64%"],
        7: ["embedder", "reranker", "Recall"],
        8: ["CP2", "CP6", "Observability"],
        9: ["RAG_PROFILE", "trace", "Telegram"],
        10: ["chunking", "protocol", "recall"],
    }

    results: list[dict] = []
    for png in sorted(preview_dir.glob("slide_*.png")):
        idx = int(png.stem.split("_")[1])
        img = Image.open(png)
        text = pytesseract.image_to_string(img, lang="rus+eng")
        clean = re.sub(r"\s+", " ", text).strip()
        char_count = len(clean)
        words = len(clean.split())
        missing = [kw for kw in expectations.get(idx, []) if kw.lower() not in clean.lower()]
        # Heuristic: too little OCR text on content slides (not title-only)
        too_sparse = idx > 1 and words < 20
        ok = char_count > 40 and not missing and not too_sparse
        results.append(
            {
                "slide": idx,
                "file": str(png),
                "ocr_chars": char_count,
                "ocr_words": words,
                "missing_keywords": missing,
                "too_sparse": too_sparse,
                "ok": ok,
                "snippet": clean[:120],
            }
        )
    return results


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
    try:
        previews = export_slide_previews(path)
        print(f"Exported {len(previews)} preview PNGs to {PREVIEW}")
        qa = ocr_qa_previews(PREVIEW)
        failed = [r for r in qa if not r["ok"]]
        print("\n=== OCR QA ===")
        for r in qa:
            status = "OK" if r["ok"] else "FAIL"
            extra = ""
            if r["missing_keywords"]:
                extra += f" missing={r['missing_keywords']}"
            if r["too_sparse"]:
                extra += " sparse"
            print(f"  slide {r['slide']:02d}: {status} ({r['ocr_words']} words){extra}")
        if failed:
            print(f"\nOCR QA: {len(failed)} slide(s) flagged (review previews manually)")
        else:
            print("\nOCR QA: all slides passed")
    except Exception as exc:
        print(f"Preview/OCR skipped: {exc}")
        sys.exit(2)
